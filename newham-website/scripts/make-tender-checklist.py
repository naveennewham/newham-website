"""Builds public/downloads/mep-tender-submission-checklist.xlsx (the download offered on
/guides/mep-tender-submission-checklist). Content follows that guide. Run:
python3 scripts/make-tender-checklist.py"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / 'public' / 'downloads' / 'mep-tender-submission-checklist.xlsx'
BRAND = '1F5F72'
HEAD, HEAD_FILL = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor=BRAND)
INPUT = PatternFill('solid', fgColor='FFF8E1')
DONE = PatternFill('solid', fgColor='E3F4E8')
THIN = Side(style='thin', color='C9D3D8')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')

STEPS = [
    ('Register the tender', 'Record the closing date, query deadline, site visit date and submission method. Set internal deadlines for take-off, pricing and review.'),
    ('Check the documents are complete', 'List every drawing, specification section, schedule and form; compare with the drawing register and query anything missing.'),
    ('Build a scope matrix by trade', 'Fire protection, ACMV, electrical, ELV, plumbing and sanitary: who supplies and installs each item, and anything by others.'),
    ('Measure quantities', 'Material and labour take-off, organised by building, floor and area.'),
    ('Compile the BOQ', "Use the client's BOQ format if given; keep descriptions aligned with the specification."),
    ('Build up the rates', 'Material and installation rates from price lists and supplier quotes, with the build-up shown.'),
    ('Get and compare quotes', 'Subcontractor and supplier quotes side by side on the same scope; note exclusions.'),
    ('Raise tender queries', 'RFIs for conflicts and anything unclear; track each query and answer.'),
    ('Review every addendum', 'Update quantities, specifications and conditions; record that each addendum is priced.'),
    ('Write qualifications and exclusions', 'State assumptions clearly; do not use them to cover scope you did not measure.'),
    ('Compile the submission', 'Priced BOQ, forms, schedules and technical submittals in the order the tender asks.'),
    ('Check and sign off', 'Run the compliance checklist; estimator or commercial manager signs off.'),
    ('File everything', 'Submitted set, query log and rate build-ups under document control for award, variations and claims.'),
]
CHECKS = [
    'Every document in the tender drawing register was received and reviewed',
    'All addenda are listed, priced and acknowledged in the form required',
    "The BOQ follows the client's format, item numbering and units",
    'No BOQ item is left blank unless the tender allows it, and blanks are explained',
    'Rate build-ups are saved and traceable to price lists or quotes',
    'Subcontractor and supplier quotes are compared on the same scope',
    'Every tender query has an answer, or the assumption made is stated',
    'Qualifications and exclusions match the pricing',
    'Totals and carried-forward figures have been checked',
    'Forms are completed and signed by the authorised person',
    'Technical submittals requested in the tender are included',
    'The submission method and closing date are confirmed',
]
TIMELINE = [
    'Query deadline (usually well before closing)', 'Subcontractor and supplier quotes due', 'Take-off complete',
    'BOQ compiled and rates built up', 'Allowance for late addenda', 'Compilation and printing / uploading',
    'Final review and sign-off by the person signing the forms', 'Tender closing date',
]
STATUS = '"Not started,In progress,Done,N/A"'


def sheet(wb, title, heading, cols, widths, rows, first_title=False):
    ws = wb.active if first_title else wb.create_sheet()
    ws.title = title
    ws['A1'] = heading
    ws['A1'].font = Font(bold=True, size=14, color=BRAND)
    for i, (c, w) in enumerate(zip(cols, widths), start=1):
        cell = ws.cell(row=3, column=i, value=c)
        cell.font, cell.fill, cell.border, cell.alignment = HEAD, HEAD_FILL, BOX, WRAP
        ws.column_dimensions[chr(64 + i)].width = w
    for r, values in enumerate(rows, start=4):
        for i, v in enumerate(values, start=1):
            cell = ws.cell(row=r, column=i, value=v)
            cell.border, cell.alignment = BOX, WRAP
    ws.freeze_panes = 'A4'
    return ws


def inputs(ws, cols, last, status_col):
    for r in range(4, last + 1):
        for c in cols:
            ws.cell(row=r, column=c).fill = INPUT
    dv = DataValidation(type='list', formula1=STATUS, allow_blank=True)
    ws.add_data_validation(dv)
    letter = chr(64 + status_col)
    dv.add(f'{letter}4:{letter}{last}')
    ws.conditional_formatting.add(f'A4:{chr(64 + ws.max_column)}{last}', FormulaRule(formula=[f'${letter}4="Done"'], fill=DONE))


wb = Workbook()
info = wb.active
info.title = 'Tender'
info['A1'] = 'MEP tender submission checklist'
info['A1'].font = Font(bold=True, size=16, color=BRAND)
info.column_dimensions['A'].width = 34
info.column_dimensions['B'].width = 48
for i, label in enumerate(['Project / tender', 'Client', 'Trades in scope', 'Closing date', 'Query deadline', 'Site visit', 'Submission method', 'Person signing the forms'], start=3):
    info.cell(row=i, column=1, value=label).font = Font(bold=True)
    c = info.cell(row=i, column=2)
    c.fill, c.border = INPUT, BOX
for i, text in enumerate([
    'Yellow cells are for you to fill in. Set Status to "Done" and the row turns green.',
    'Sheets: Sequence (13 steps), Compliance (check before you submit), Timeline (work back from closing).',
    'From the guide: newhamtech.com/guides/mep-tender-submission-checklist',
    'Need take-offs, a priced BOQ or tender support? newhamtech.com/boq-estimation-services · admin@newhamtech.com · +65 9880 4197',
], start=13):
    info.cell(row=i, column=1, value=text).font = Font(size=9, color='5B6B73')

seq = sheet(wb, 'Sequence', 'The tender sequence, step by step', ['#', 'Step', 'What to do', 'Owner', 'Due', 'Status', 'Notes'],
            [5, 28, 60, 16, 12, 14, 30], [(i, s, d, None, None, None, None) for i, (s, d) in enumerate(STEPS, start=1)])
inputs(seq, (4, 5, 6, 7), 3 + len(STEPS), 6)
chk = sheet(wb, 'Compliance', 'Compliance checklist before you submit', ['#', 'Check', 'Status', 'Checked by', 'Notes'],
            [5, 70, 14, 18, 34], [(i, c, None, None, None) for i, c in enumerate(CHECKS, start=1)])
inputs(chk, (3, 4, 5), 3 + len(CHECKS), 3)
tl = sheet(wb, 'Timeline', 'Planning the internal timeline (work back from the closing date)', ['#', 'Milestone', 'Date', 'Owner', 'Status'],
           [5, 56, 14, 18, 14], [(i, m, None, None, None) for i, m in enumerate(TIMELINE, start=1)])
inputs(tl, (3, 4, 5), 3 + len(TIMELINE), 5)
for r in range(4, 4 + len(TIMELINE)):
    tl.cell(row=r, column=3).number_format = 'dd mmm yyyy'

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
