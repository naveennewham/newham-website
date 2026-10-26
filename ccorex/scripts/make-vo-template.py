"""Builds public/downloads/variation-order-template-singapore.xlsx (the download offered on
/resources/variation-order-template-singapore). Layout follows that guide: VO form with pricing,
VO register, notes. Run: python3 scripts/make-vo-template.py"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / 'public' / 'downloads' / 'variation-order-template-singapore.xlsx'
BRAND = '1F5F72'
HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor=BRAND)
INPUT_FILL = PatternFill('solid', fgColor='FFF8E1')  # yellow cells are for the user to fill
TOTAL_FILL = PatternFill('solid', fgColor='E6F1F4')
THIN = Side(style='thin', color='C9D3D8')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0.00;(#,##0.00);"-"'
QTY = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%'
LINES = 25          # pricing lines on the VO form
REGISTER_ROWS = 40  # rows in the VO register
GREY = Font(size=9, color='5B6B73')


def header_row(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill, c.border = HEAD, HEAD_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical='center')
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + i)].width = w


def total_cell(ws, row, col, value, fmt=MONEY):
    c = ws.cell(row=row, column=col, value=value)
    c.number_format, c.font, c.fill, c.border = fmt, Font(bold=True), TOTAL_FILL, BOX
    return c


wb = Workbook()

# --- VO form ---
f = wb.active
f.title = 'VO form'
f['A1'] = 'VARIATION ORDER'
f['A1'].font = Font(bold=True, size=16, color=BRAND)
header_fields = [
    'Project (title and site)', 'Contract / subcontract reference', 'VO number', 'Instruction reference and date',
    'Instructed by (name and role)', 'Description of change', 'Reason / cause', 'Drawings and documents', 'Time impact',
]
for i, label in enumerate(header_fields, start=3):
    f.cell(row=i, column=2, value=label).font = Font(bold=True)
    f.merge_cells(start_row=i, start_column=3, end_row=i, end_column=6)
    for col in range(3, 7):
        c = f.cell(row=i, column=col)
        c.fill, c.border = INPUT_FILL, BOX
f['C5'] = 'VO 001'

first = 14
header_row(f, first - 1, ['Item', 'Description', 'Qty (negative = omission)', 'Unit', 'Rate (S$)', 'Amount (S$)'],
           [8, 46, 14, 9, 14, 18])
for r in range(first, first + LINES):
    for col in range(1, 7):
        f.cell(row=r, column=col).border = BOX
    for col in (1, 2, 3, 4, 5):
        f.cell(row=r, column=col).fill = INPUT_FILL
    f.cell(row=r, column=3).number_format = QTY
    f.cell(row=r, column=5).number_format = MONEY
    f.cell(row=r, column=6, value=f'=IF(OR(C{r}="",E{r}=""),"",C{r}*E{r})').number_format = MONEY

last = first + LINES - 1
net, mk_rate, mk, value, gst_rate, gst, total = range(last + 1, last + 8)
f.cell(row=net, column=2, value='Net value of work (additions less omissions)').font = Font(bold=True)
total_cell(f, net, 6, f'=SUM(F{first}:F{last})')
f.cell(row=mk_rate, column=2, value='Markup / preliminaries rate (only if the contract allows)')
c = f.cell(row=mk_rate, column=6, value=0)
c.fill, c.border, c.number_format = INPUT_FILL, BOX, PCT
f.cell(row=mk, column=2, value='Markup / preliminaries')
c = f.cell(row=mk, column=6, value=f'=F{net}*F{mk_rate}')
c.number_format, c.border = MONEY, BOX
f.cell(row=value, column=2, value='Variation value (excl. GST)').font = Font(bold=True)
total_cell(f, value, 6, f'=F{net}+F{mk}')
f.cell(row=gst_rate, column=2, value='GST rate (blank if not GST-registered)')
c = f.cell(row=gst_rate, column=6, value=0.09)
c.fill, c.border, c.number_format = INPUT_FILL, BOX, PCT
f.cell(row=gst, column=2, value='GST')
c = f.cell(row=gst, column=6, value=f'=IF(F{gst_rate}="",0,F{value}*F{gst_rate})')
c.number_format, c.border = MONEY, BOX
f.cell(row=total, column=2, value='Total including GST').font = Font(bold=True)
total_cell(f, total, 6, f'=F{value}+F{gst}')

sign = total + 2
for i, label in enumerate(['Submitted by (name, signature, date)', 'Agreed by (name, signature, date)']):
    f.cell(row=sign + i, column=2, value=label).font = Font(bold=True)
    f.merge_cells(start_row=sign + i, start_column=3, end_row=sign + i, end_column=6)
    for col in range(3, 7):
        c = f.cell(row=sign + i, column=col)
        c.fill, c.border = INPUT_FILL, BOX
for i, text in enumerate([
    'Yellow cells are inputs. Enter omissions as negative quantities; amounts are calculated.',
    'Rates, markup and notice periods come from your contract; check them before you submit the VO.',
    'Template by CCorex (ccorex.com). General information, not legal advice.',
]):
    c = f.cell(row=sign + 3 + i, column=2, value=text)
    c.font = GREY
    f.merge_cells(start_row=sign + 3 + i, start_column=2, end_row=sign + 3 + i, end_column=6)
f.freeze_panes = f'A{first}'
f.sheet_view.showGridLines = False

# --- VO register ---
g = wb.create_sheet('VO register')
g['A1'] = 'Variation order register'
g['A1'].font = Font(bold=True, size=13, color=BRAND)
header_row(g, 3, ['VO number', 'Description', 'Instruction date', 'Submitted date', 'Status', 'Submitted value (S$)',
                  'Agreed value (S$)', 'Included in claim no.', 'Remarks'],
           [11, 40, 14, 14, 13, 18, 18, 14, 32])
r1 = 4
rN = r1 + REGISTER_ROWS - 1
for r in range(r1, rN + 1):
    for col in range(1, 10):
        c = g.cell(row=r, column=col)
        c.border, c.fill = BOX, INPUT_FILL
    g.cell(row=r, column=3).number_format = 'dd-mmm-yyyy'
    g.cell(row=r, column=4).number_format = 'dd-mmm-yyyy'
    g.cell(row=r, column=6).number_format = MONEY
    g.cell(row=r, column=7).number_format = MONEY
status = DataValidation(type='list', formula1='"Submitted,Agreed,Rejected,Disputed"', allow_blank=True)
g.add_data_validation(status)
status.add(f'E{r1}:E{rN}')

t = rN + 1
g.cell(row=t, column=2, value='Total').font = Font(bold=True)
total_cell(g, t, 6, f'=SUM(F{r1}:F{rN})')
total_cell(g, t, 7, f'=SUM(G{r1}:G{rN})')
summary = [
    ('Submitted value awaiting agreement', f'=SUMIF(E{r1}:E{rN},"Submitted",F{r1}:F{rN})'),
    ('Submitted value under dispute', f'=SUMIF(E{r1}:E{rN},"Disputed",F{r1}:F{rN})'),
    ('Agreed value', f'=SUMIF(E{r1}:E{rN},"Agreed",G{r1}:G{rN})'),
    ('Number of VOs not yet agreed', f'=COUNTIF(E{r1}:E{rN},"Submitted")+COUNTIF(E{r1}:E{rN},"Disputed")'),
]
for i, (label, formula) in enumerate(summary, start=t + 2):
    g.cell(row=i, column=2, value=label).font = Font(bold=True)
    c = g.cell(row=i, column=6, value=formula)
    c.border = BOX
    c.number_format = '0' if label.startswith('Number') else MONEY
g.freeze_panes = 'A4'

# --- Notes ---
n = wb.create_sheet('Notes')
n.column_dimensions['A'].width = 100
n['A1'] = 'How to use this template'
n['A1'].font = Font(bold=True, size=13, color=BRAND)
for i, text in enumerate([
    '1. Copy the "VO form" sheet for each variation, or keep one file per VO. Number VOs in one sequence per project.',
    '2. Fill the header from the written instruction: reference, date and who instructed it.',
    '3. Price each line. Use contract rates first, then pro-rata rates, then quotations or daywork, as your contract allows.',
    '4. Enter omissions as negative quantities so the net value is clear.',
    '5. Add markup or preliminaries only if your contract allows it and your BOQ rates do not already include them.',
    '6. Record every VO on the "VO register" sheet and update its status as it is agreed, rejected or disputed.',
    '7. Before each progress claim, check the register and claim the VOs your contract allows, on a separate variations sheet.',
    '',
    'Your contract (for example the SIA conditions, PSSCOC or REDAS conditions, or your subcontract) sets who may instruct',
    'variations, how they must be confirmed, notice periods and valuation rules. Check it for each project.',
    'This template is general information, not legal advice. Guide: ccorex.com/resources/variation-order-template-singapore',
], start=3):
    n.cell(row=i, column=1, value=text).alignment = Alignment(wrap_text=True)

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
