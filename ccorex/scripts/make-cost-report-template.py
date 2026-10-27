"""Builds public/downloads/construction-cost-report-template.xlsx (the download offered on
/resources/construction-cost-report-template). Layout follows that guide: monthly cost report by
cost code with revenue and margin, a monthly trend sheet, notes. Run: python3 scripts/make-cost-report-template.py"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = Path(__file__).resolve().parent.parent / 'public' / 'downloads' / 'construction-cost-report-template.xlsx'
BRAND = '1F5F72'
HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor=BRAND)
INPUT_FILL = PatternFill('solid', fgColor='FFF8E1')  # yellow cells are for the user to fill
TOTAL_FILL = PatternFill('solid', fgColor='E6F1F4')
RED = Font(color='B42318')
THIN = Side(style='thin', color='C9D3D8')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0;(#,##0);"-"'
PCT = '0.0%'
GREY = Font(size=9, color='5B6B73')
CODES = ['Materials', 'Subcontract labour', 'Direct labour', 'Equipment and testing', 'Preliminaries', 'Overhead']
ROWS = 20  # cost code rows
MONTHS = 24


def header_row(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill, c.border = HEAD, HEAD_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical='center')
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + i)].width = w


def total_cell(ws, row, col, value, fmt=MONEY):
    c = ws.cell(row=row, column=col, value=value)
    c.number_format, c.font, c.fill, c.border = fmt, Font(bold=True), TOTAL_FILL, BOX
    return c


def input_cell(ws, row, col, value=None, fmt=MONEY):
    c = ws.cell(row=row, column=col, value=value)
    c.fill, c.border, c.number_format = INPUT_FILL, BOX, fmt
    return c


wb = Workbook()

# --- Cost report ---
s = wb.active
s.title = 'Cost report'
s['A1'] = 'MONTHLY PROJECT COST REPORT'
s['A1'].font = Font(bold=True, size=16, color=BRAND)
for i, label in enumerate(['Project', 'Report month', 'Prepared by', 'Currency'], start=3):
    s.cell(row=i, column=1, value=label).font = Font(bold=True)
    s.merge_cells(start_row=i, start_column=2, end_row=i, end_column=4)
    for col in range(2, 5):
        c = s.cell(row=i, column=col)
        c.fill, c.border = INPUT_FILL, BOX
s['B6'] = 'SGD, before GST'

# Revenue block
s['A8'] = 'Revenue'
s['A8'].font = Font(bold=True, size=12, color=BRAND)
rev = [('Contract value', None), ('Approved variations', None), ('Expected variations (not yet agreed)', None),
       ('Claimed to date', None), ('Certified to date', None)]
for i, (label, _) in enumerate(rev, start=9):
    s.cell(row=i, column=1, value=label)
    input_cell(s, i, 2)
s.cell(row=14, column=1, value='Forecast final revenue').font = Font(bold=True)
total_cell(s, 14, 2, '=B9+B10+B11')

# Cost table
first = 18
s['A17'] = 'Cost by cost code'
s['A17'].font = Font(bold=True, size=12, color=BRAND)
header_row(s, first, ['Cost code', 'Budget', 'Committed (POs and subcontracts)', 'Actual to date',
                      'Cost to complete (PM estimate)', 'Forecast final cost', 'Variance (budget less forecast)',
                      'Variance %', 'Reason and action'],
           [30, 14, 16, 14, 16, 16, 16, 11, 44])
r1 = first + 1
rN = r1 + ROWS - 1
for r in range(r1, rN + 1):
    code = CODES[r - r1] if r - r1 < len(CODES) else None
    input_cell(s, r, 1, code, fmt='General')
    for col in (2, 3, 4, 5):
        input_cell(s, r, col)
    c = s.cell(row=r, column=6, value=f'=IF(AND(D{r}="",E{r}=""),"",N(D{r})+N(E{r}))')
    c.number_format, c.border = MONEY, BOX
    c = s.cell(row=r, column=7, value=f'=IF(F{r}="","",N(B{r})-F{r})')
    c.number_format, c.border = MONEY, BOX
    c = s.cell(row=r, column=8, value=f'=IF(OR(F{r}="",N(B{r})=0),"",G{r}/B{r})')
    c.number_format, c.border = PCT, BOX
    input_cell(s, r, 9, fmt='General').alignment = Alignment(wrap_text=True)
s.conditional_formatting.add(f'G{r1}:H{rN}', CellIsRule(operator='lessThan', formula=['0'], font=RED))
t = rN + 1
s.cell(row=t, column=1, value='Total').font = Font(bold=True)
for col, letter in zip(range(2, 8), 'BCDEFG'):
    total_cell(s, t, col, f'=SUM({letter}{r1}:{letter}{rN})')
total_cell(s, t, 8, f'=IF(B{t}=0,"",G{t}/B{t})', fmt=PCT)

# Margin block
m = t + 2
s.cell(row=m, column=1, value='Margin').font = Font(bold=True, size=12, color=BRAND)
margin = [
    ('Budget margin (contract value + approved variations − budget)', f'=B9+B10-B{t}'),
    ('Forecast margin (forecast final revenue − forecast final cost)', f'=B14-F{t}'),
    ('Forecast margin %', f'=IF(B14=0,"",(B14-F{t})/B14)'),
    ('Committed but not yet invoiced (check cost to complete covers it)', f'=MAX(0,C{t}-D{t})'),
]
for i, (label, formula) in enumerate(margin, start=m + 1):
    s.cell(row=i, column=1, value=label)
    total_cell(s, i, 2, formula, fmt=PCT if label.endswith('%') else MONEY)
for i, text in enumerate([
    'Yellow cells are inputs. Forecast final cost = actual to date + cost to complete. Negative variance (red) means over budget.',
    'Cost to complete is the PM\'s estimate of what is still needed, including commitments not yet invoiced. Do not set it to budget minus actual.',
    'Template by CCorex (ccorex.com). General information, not accounting advice.',
], start=m + 6):
    c = s.cell(row=i, column=1, value=text)
    c.font = GREY
    s.merge_cells(start_row=i, start_column=1, end_row=i, end_column=9)
s.freeze_panes = f'B{r1}'
s.sheet_view.showGridLines = False

# --- Monthly trend ---
g = wb.create_sheet('Monthly trend')
g['A1'] = 'Forecast final cost and margin by month'
g['A1'].font = Font(bold=True, size=13, color=BRAND)
header_row(g, 3, ['Month', 'Budget', 'Actual to date', 'Forecast final cost', 'Forecast final revenue',
                  'Forecast margin', 'Forecast margin %', 'Comment'],
           [12, 14, 14, 16, 16, 14, 12, 44])
for r in range(4, 4 + MONTHS):
    input_cell(g, r, 1, fmt='mmm-yyyy')
    for col in (2, 3, 4, 5):
        input_cell(g, r, col)
    c = g.cell(row=r, column=6, value=f'=IF(OR(D{r}="",E{r}=""),"",E{r}-D{r})')
    c.number_format, c.border = MONEY, BOX
    c = g.cell(row=r, column=7, value=f'=IF(OR(F{r}="",N(E{r})=0),"",F{r}/E{r})')
    c.number_format, c.border = PCT, BOX
    input_cell(g, r, 8, fmt='General')
g.cell(row=4 + MONTHS + 1, column=1, value='Copy the totals from the "Cost report" sheet here at each month-end, so the trend is kept.').font = GREY
g.freeze_panes = 'A4'

# --- Notes ---
n = wb.create_sheet('Notes')
n.column_dimensions['A'].width = 100
n['A1'] = 'How to use this template'
n['A1'].font = Font(bold=True, size=13, color=BRAND)
for i, text in enumerate([
    '1. Use the same cost codes as your approved budget, so budget and actuals line up row by row.',
    '2. Close the month first: post matched invoices and allocate labour and overhead, so "Actual to date" is complete.',
    '3. "Committed" is every approved PO and subcontract value, invoiced or not. Remove cancelled orders.',
    '4. Ask the PM for "Cost to complete" code by code, in figures, with a reason for any change from last month.',
    '5. Explain every variance above your threshold in the "Reason and action" column.',
    '6. Fill the revenue block, including approved variations and, separately, variations you expect to be agreed.',
    '7. Copy the month\'s totals to the "Monthly trend" sheet and keep each month\'s report.',
    '',
    'Figures are before GST. This template is general information, not accounting advice.',
    'Guide: ccorex.com/resources/construction-cost-report-template',
], start=3):
    n.cell(row=i, column=1, value=text).alignment = Alignment(wrap_text=True)

OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
