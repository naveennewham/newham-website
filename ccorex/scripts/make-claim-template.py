"""Builds public/downloads/progress-claim-template-singapore.xlsx (the download offered on
/resources/progress-claim-template-singapore). Layout follows that guide: header, valuation
sheet, variations sheet, summary. Run: python3 scripts/make-claim-template.py"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / 'public' / 'downloads' / 'progress-claim-template-singapore.xlsx'
BRAND = '1F5F72'
HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor=BRAND)
INPUT_FILL = PatternFill('solid', fgColor='FFF8E1')  # yellow cells are for the user to fill
TOTAL_FILL = PatternFill('solid', fgColor='E6F1F4')
THIN = Side(style='thin', color='C9D3D8')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0.00;(#,##0.00);"-"'
PCT = '0.0%'
ROWS = 30  # item rows on each sheet


def header_row(ws, row, labels, widths=None):
    for i, label in enumerate(labels, start=1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill, c.border = HEAD, HEAD_FILL, BOX
        c.alignment = Alignment(wrap_text=True, vertical='center')
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[chr(64 + i)].width = w


def item_table(ws, first, extra_cols):
    """Item rows: A item, B description, C contract amount, D % to date, E value to date (=C*D),
    F previously claimed, G this period (=E-F), then extra columns. Returns the totals row."""
    for r in range(first, first + ROWS):
        for col in range(1, 8 + len(extra_cols)):
            ws.cell(row=r, column=col).border = BOX
        for col in (1, 2, 3, 4, 6, *range(8, 8 + len(extra_cols))):
            ws.cell(row=r, column=col).fill = INPUT_FILL
        ws.cell(row=r, column=3).number_format = MONEY
        ws.cell(row=r, column=4).number_format = PCT
        ws.cell(row=r, column=5, value=f'=IF(C{r}="","",C{r}*D{r})').number_format = MONEY
        ws.cell(row=r, column=6).number_format = MONEY
        ws.cell(row=r, column=7, value=f'=IF(E{r}="","",E{r}-F{r})').number_format = MONEY
    total = first + ROWS
    ws.cell(row=total, column=2, value='Total').font = Font(bold=True)
    for col, letter in ((3, 'C'), (5, 'E'), (6, 'F'), (7, 'G')):
        c = ws.cell(row=total, column=col, value=f'=SUM({letter}{first}:{letter}{total - 1})')
        c.number_format, c.font, c.fill, c.border = MONEY, Font(bold=True), TOTAL_FILL, BOX
    pct = ws.cell(row=total + 2, column=2, value='Tip: enter % complete (e.g. 40%), or overwrite column E with quantity × rate.')
    pct.font = Font(italic=True, color='5B6B73')
    return total


wb = Workbook()

# --- Summary (with the header) ---
s = wb.active
s.title = 'Summary'
s.column_dimensions['A'].width = 6
s.column_dimensions['B'].width = 46
s.column_dimensions['C'].width = 22
s['A1'] = 'PROGRESS CLAIM'
s['A1'].font = Font(bold=True, size=16, color=BRAND)
header_fields = [
    ('Claimant (company name and UEN)', ''), ('Respondent', ''), ('Project (title and site)', ''),
    ('Contract / subcontract reference', ''), ('Claim number', 'PC 01'), ('Claim period (work up to)', ''),
    ('Date of claim', ''),
]
for i, (label, value) in enumerate(header_fields, start=3):
    s.cell(row=i, column=2, value=label).font = Font(bold=True)
    c = s.cell(row=i, column=3, value=value)
    c.fill, c.border = INPUT_FILL, BOX
sop = s.cell(row=11, column=2, value='This is a payment claim made under the Building and Construction Industry Security of Payment Act 2004.')
sop.font = Font(italic=True)
s.merge_cells('B11:C11')
s.cell(row=12, column=2, value='(Optional statement. Delete this line if you do not want to use it.)').font = Font(size=9, color='5B6B73')

inputs = 14
s.cell(row=inputs, column=2, value='Inputs').font = Font(bold=True, color=BRAND)
for i, (label, value, fmt) in enumerate([
    ('Contract sum (S$)', None, MONEY), ('Retention rate', 0.05, PCT), ('Retention limit (S$, blank = none)', None, MONEY),
    ('Materials on site (S$, if the contract allows)', 0, MONEY), ('Previous certified amounts, excl. GST (S$)', 0, MONEY),
    ('GST rate (blank if not GST-registered)', 0.09, PCT),
], start=inputs + 1):
    s.cell(row=i, column=2, value=label)
    c = s.cell(row=i, column=3, value=value)
    c.fill, c.border, c.number_format = INPUT_FILL, BOX, fmt
# Named rows for readability in formulas
CS, RR, RL, MOS, PREV, GST = (f'C{inputs + n}' for n in range(1, 7))

lines_at = inputs + 8
header_row(s, lines_at, ['Line', 'Item', 'Amount (S$)'])
lines = [
    ('Contract sum', f'={CS}'),
    ('Contract work done to date', f"='Valuation'!E{4 + ROWS}"),
    ('Variations done to date', f"='Variations'!E{4 + ROWS}"),
    ('Materials on site', f'={MOS}'),
    ('Gross value to date (2 + 3 + 4)', None),
    ('Less retention (rate × line 5, up to the limit)', None),
    ('Net value to date (5 − 6)', None),
    ('Less previous certified amounts', f'={PREV}'),
    ('Amount claimed this period (7 − 8)', None),
    ('GST', None),
    ('Total amount claimed this period (9 + 10)', None),
]
r0 = lines_at + 1
row = {n: r0 + n - 1 for n in range(1, 12)}
formulas = {
    5: f'=C{row[2]}+C{row[3]}+C{row[4]}',
    6: f'=IF({RL}="",C{row[5]}*{RR},MIN(C{row[5]}*{RR},{RL}))',
    7: f'=C{row[5]}-C{row[6]}',
    9: f'=C{row[7]}-C{row[8]}',
    10: f'=IF({GST}="",0,C{row[9]}*{GST})',
    11: f'=C{row[9]}+C{row[10]}',
}
for n, (label, f) in enumerate(lines, start=1):
    r = row[n]
    s.cell(row=r, column=1, value=n).border = BOX
    s.cell(row=r, column=2, value=label).border = BOX
    c = s.cell(row=r, column=3, value=formulas.get(n, f))
    c.number_format, c.border = MONEY, BOX
    if n in (5, 7, 9, 11):
        for col in (1, 2, 3):
            s.cell(row=r, column=col).font = Font(bold=True)
            s.cell(row=r, column=col).fill = TOTAL_FILL
note_r = row[11] + 2
for i, text in enumerate([
    'Yellow cells are inputs. Everything else is calculated.',
    'Retention, payment terms and GST treatment come from your contract; check them before you send the claim.',
    'This template is general information, not legal advice. Check the current SOP Act and take legal advice on disputes.',
    'Template by CCorex (ccorex.com). DeepTrack prepares numbered progress claims with retention, 9% GST and a PDF certificate.',
]):
    c = s.cell(row=note_r + i, column=2, value=text)
    c.font = Font(size=9, color='5B6B73')
    s.merge_cells(start_row=note_r + i, start_column=2, end_row=note_r + i, end_column=3)

# --- Valuation sheet ---
v = wb.create_sheet('Valuation')
v['A1'] = 'Valuation sheet (contract work)'
v['A1'].font = Font(bold=True, size=13, color=BRAND)
header_row(v, 3, ['Item', 'Description', 'Contract amount (S$)', '% to date', 'Value to date (S$)', 'Previously claimed (S$)', 'This period (S$)', 'Remarks'],
           [8, 40, 18, 11, 18, 18, 16, 30])
item_table(v, 4, ['Remarks'])

# --- Variations sheet ---
vo = wb.create_sheet('Variations')
vo['A1'] = 'Variations'
vo['A1'].font = Font(bold=True, size=13, color=BRAND)
header_row(vo, 3, ['VO ref', 'Description', 'VO amount (S$)', '% to date', 'Value to date (S$)', 'Previously claimed (S$)', 'This period (S$)', 'Status', 'Remarks'],
           [10, 38, 16, 11, 18, 18, 16, 14, 26])
item_table(vo, 4, ['Status', 'Remarks'])
status = DataValidation(type='list', formula1='"Instructed,Submitted,Agreed,Disputed"', allow_blank=True)
vo.add_data_validation(status)
status.add(f'H4:H{3 + ROWS}')

for ws in (v, vo):
    ws.freeze_panes = 'A4'
s.sheet_view.showGridLines = False
OUT.parent.mkdir(parents=True, exist_ok=True)
wb.save(OUT)
print(OUT)
